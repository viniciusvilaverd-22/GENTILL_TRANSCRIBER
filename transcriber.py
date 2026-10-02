from __future__ import annotations

import argparse
import json
import sys
import time
from dataclasses import asdict, dataclass, replace
from pathlib import Path
from typing import Callable, Iterable, Sequence


SUPPORTED_INPUTS = {
    ".wav", ".mp3", ".m4a", ".aac", ".flac", ".ogg", ".opus",
    ".mp4", ".mkv", ".mov", ".webm", ".avi",
}

ProgressCallback = Callable[[float, float | None, "Segment"], None]


class TranscriptionCancelled(RuntimeError):
    """Raised when a caller cancels a running transcription."""


@dataclass(frozen=True)
class Segment:
    start: float
    end: float
    text: str
    speaker: str | None = None


def _timestamp_srt(seconds: float) -> str:
    total_ms = max(0, int(round(seconds * 1000)))
    hours, rem = divmod(total_ms, 3_600_000)
    minutes, rem = divmod(rem, 60_000)
    secs, millis = divmod(rem, 1000)
    return f"{hours:02}:{minutes:02}:{secs:02},{millis:03}"


def _timestamp_vtt(seconds: float) -> str:
    return _timestamp_srt(seconds).replace(",", ".")


def _segment_text(segment: Segment) -> str:
    text = segment.text.strip()
    if segment.speaker:
        return f"[{segment.speaker}] {text}"
    return text


def render_txt(segments: Sequence[Segment]) -> str:
    return "\n".join(_segment_text(s) for s in segments if s.text.strip()).strip() + "\n"


def render_srt(segments: Sequence[Segment]) -> str:
    blocks = []
    for index, segment in enumerate(segments, start=1):
        blocks.append(
            f"{index}\n"
            f"{_timestamp_srt(segment.start)} --> {_timestamp_srt(segment.end)}\n"
            f"{_segment_text(segment)}\n"
        )
    return "\n".join(blocks)


def render_vtt(segments: Sequence[Segment]) -> str:
    blocks = ["WEBVTT\n"]
    for segment in segments:
        blocks.append(
            f"{_timestamp_vtt(segment.start)} --> {_timestamp_vtt(segment.end)}\n"
            f"{_segment_text(segment)}\n"
        )
    return "\n".join(blocks)


def write_docx(output_path: Path, segments: Sequence[Segment], metadata: dict) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Exportação DOCX requer o pacote python-docx.") from exc

    document = Document()
    document.add_heading("Gentill Transcriber", level=0)
    document.add_paragraph(f"Arquivo: {metadata.get('input', '')}")
    document.add_paragraph(f"Idioma: {metadata.get('language', 'auto')}")
    document.add_paragraph(f"Modelo: {metadata.get('model', '')}")
    document.add_paragraph("")

    for segment in segments:
        stamp = f"{_timestamp_vtt(segment.start)} → {_timestamp_vtt(segment.end)}"
        prefix = f"{segment.speaker} · " if segment.speaker else ""
        paragraph = document.add_paragraph()
        paragraph.add_run(prefix + stamp + "\n").bold = True
        paragraph.add_run(segment.text.strip())

    document.save(output_path)


def write_outputs(
    output_dir: Path,
    stem: str,
    segments: Sequence[Segment],
    metadata: dict,
    formats: set[str] | None = None,
) -> list[Path]:
    """Write requested outputs. Defaults preserve the original TXT/SRT/VTT/JSON contract."""
    selected = formats or {"txt", "srt", "vtt", "json"}
    output_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    if "txt" in selected:
        path = output_dir / f"{stem}.txt"
        path.write_text(render_txt(segments), encoding="utf-8")
        written.append(path)
    if "srt" in selected:
        path = output_dir / f"{stem}.srt"
        path.write_text(render_srt(segments), encoding="utf-8")
        written.append(path)
    if "vtt" in selected:
        path = output_dir / f"{stem}.vtt"
        path.write_text(render_vtt(segments), encoding="utf-8")
        written.append(path)
    if "json" in selected:
        path = output_dir / f"{stem}.json"
        payload = {
            "metadata": metadata,
            "segments": [asdict(segment) for segment in segments],
        }
        path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        written.append(path)
    if "docx" in selected:
        path = output_dir / f"{stem}.docx"
        write_docx(path, segments, metadata)
        written.append(path)

    return written


def load_model(model_path: Path, device: str, compute_type: str):
    if not model_path.exists():
        raise FileNotFoundError(
            f"Modelo local não encontrado: {model_path}. "
            "Downloads implícitos ficam desativados durante a transcrição."
        )

    try:
        from faster_whisper import WhisperModel
    except ImportError as exc:
        raise RuntimeError("Dependência faster-whisper não instalada.") from exc

    return WhisperModel(
        str(model_path),
        device=device,
        compute_type=compute_type,
        local_files_only=True,
    )


def _event_is_set(event: object | None) -> bool:
    return bool(event is not None and getattr(event, "is_set", lambda: False)())


def _wait_if_paused(pause_event: object | None, cancel_event: object | None) -> None:
    while _event_is_set(pause_event):
        if _event_is_set(cancel_event):
            raise TranscriptionCancelled("Transcrição cancelada.")
        time.sleep(0.08)


def transcribe(
    input_path: Path,
    model_path: Path,
    language: str | None,
    device: str,
    compute_type: str,
    progress_callback: ProgressCallback | None = None,
    cancel_event: object | None = None,
    pause_event: object | None = None,
) -> tuple[list[Segment], dict]:
    if not input_path.is_file():
        raise FileNotFoundError(f"Entrada não encontrada: {input_path}")
    if input_path.suffix.lower() not in SUPPORTED_INPUTS:
        raise ValueError(f"Formato não permitido: {input_path.suffix}")
    if _event_is_set(cancel_event):
        raise TranscriptionCancelled("Transcrição cancelada.")

    model = load_model(model_path, device=device, compute_type=compute_type)
    _wait_if_paused(pause_event, cancel_event)

    raw_segments, info = model.transcribe(
        str(input_path),
        language=language,
        vad_filter=True,
        beam_size=5,
    )

    duration = getattr(info, "duration", None)
    duration_value = float(duration) if duration not in (None, 0) else None
    segments: list[Segment] = []

    for raw in raw_segments:
        if _event_is_set(cancel_event):
            raise TranscriptionCancelled("Transcrição cancelada.")
        _wait_if_paused(pause_event, cancel_event)
        text = str(raw.text).strip()
        if not text:
            continue
        segment = Segment(
            start=float(raw.start),
            end=float(raw.end),
            text=text,
        )
        segments.append(segment)
        if progress_callback:
            fraction = (
                min(1.0, max(0.0, segment.end / duration_value))
                if duration_value
                else 0.0
            )
            progress_callback(fraction, duration_value, segment)

    if progress_callback and segments:
        progress_callback(1.0, duration_value, segments[-1])

    metadata = {
        "input": str(input_path.resolve()),
        "model": str(model_path.resolve()),
        "language": getattr(info, "language", language),
        "language_probability": getattr(info, "language_probability", None),
        "duration": duration,
        "device": device,
        "compute_type": compute_type,
        "offline_policy": "local_files_only",
    }
    return segments, metadata


def apply_speaker_turns(
    segments: Sequence[Segment],
    turns: Sequence[tuple[float, float, str]],
) -> list[Segment]:
    """Assign each segment to the speaker with the greatest time overlap."""
    result: list[Segment] = []
    for segment in segments:
        best_speaker: str | None = None
        best_overlap = 0.0
        for start, end, speaker in turns:
            overlap = max(0.0, min(segment.end, end) - max(segment.start, start))
            if overlap > best_overlap:
                best_overlap = overlap
                best_speaker = speaker
        result.append(replace(segment, speaker=best_speaker))
    return result


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Gentill Transcriber — offline")
    parser.add_argument("input", type=Path)
    parser.add_argument("--model", type=Path, default=Path("models/whisper-small-ct2"))
    parser.add_argument("--output", type=Path, default=Path("output"))
    parser.add_argument("--language", default="pt", help="Use 'auto' para autodetecção.")
    parser.add_argument("--device", default="cpu", choices=["cpu", "cuda", "auto"])
    parser.add_argument("--compute-type", default="int8")
    parser.add_argument(
        "--formats",
        default="txt,srt,vtt,json",
        help="Lista separada por vírgulas: txt,srt,vtt,json,docx",
    )
    return parser


def main(argv: Iterable[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    language = None if args.language.lower() == "auto" else args.language.lower()
    formats = {item.strip().lower() for item in args.formats.split(",") if item.strip()}

    try:
        segments, metadata = transcribe(
            input_path=args.input,
            model_path=args.model,
            language=language,
            device=args.device,
            compute_type=args.compute_type,
        )
        write_outputs(args.output, args.input.stem, segments, metadata, formats=formats)
    except Exception as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        return 1

    print(f"TRANSCRIPTION=PASS segments={len(segments)} output={args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
